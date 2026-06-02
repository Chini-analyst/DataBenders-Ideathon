"""DOCX parser using python-docx."""
from __future__ import annotations


def parse_docx(file_path: str) -> str:
    """
    Parse a .docx file and return its text content.
    Includes paragraphs and table cell text.
    """
    try:
        from docx import Document

        doc = Document(file_path)
        parts = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)

        return "\n".join(parts)
    except Exception as exc:
        raise ValueError(f"Failed to parse DOCX file: {exc}") from exc
